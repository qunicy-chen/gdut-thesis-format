# -*- coding: utf-8 -*-
"""
《智能电网》课程报告 docx 生成脚本
格式依据: 广东工业大学本科生毕业设计(论文)格式规范
  A4, 页边距 上30/下25/左30/右20mm, 1.5倍行距
  章标题 三号黑体加粗(居中), 节标题 小四黑体加粗, 正文 小四宋体(西文TNR)
  表题表上五号黑体加粗, 图题图下五号宋体, 三线表
  页码从正文起 阿拉伯数字 页脚右侧 小五TNR
用法: python build_report.py [toc_data.json]   (带参数时回填目录页码)
"""
import re, sys, os, json
import docx
from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    from docx.oxml import parse_xml
except ImportError:
    from docx.oxml.parser import parse_xml

WORK = r'C:\Users\<你的用户名>\Desktop\论文项目\_work'
OUTDIR = r'C:\Users\<你的用户名>\Desktop\论文项目\交付'
FIGS = WORK + r'\figs'
DOC_PATH = OUTDIR + r'\智能配电网光储协同的电压治理系统方案设计_课程报告.docx'

W_DECL = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
M_DECL = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'

# ===================== OMML 公式构造 =====================
def mr(t):
    return f'<m:r><m:t xml:space="preserve">{t}</m:t></m:r>'

def frac(num, den):
    return f'<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>'

def ssub(e, s):
    return f'<m:sSub><m:e>{e}</m:e><m:sub>{s}</m:sub></m:sSub>'

def ssup(e, s):
    return f'<m:sSup><m:e>{e}</m:e><m:sup>{s}</m:sup></m:sSup>'

def rad(e):
    return f'<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>{e}</m:e></m:rad>'

def dparen(e):
    return f'<m:d><m:e>{e}</m:e></m:d>'

def omath(content):
    return f'<m:oMath {M_DECL} {W_DECL}>{content}</m:oMath>'

UN = ssub(mr('U'), mr('N'))
FORMULAS = {
    '2.1': omath(mr('ΔU=') + frac(mr('PR+QX'), UN)),
    '2.2': omath(ssub(mr('U'), mr('2')) + mr('≈') + ssub(mr('U'), mr('1')) + mr('−') +
                 frac(mr('PR+QX'), UN)),
    '2.3': omath(frac(mr('∂') + ssub(mr('U'), mr('2')), mr('∂P')) + mr('=−') + frac(mr('R'), UN) +
                 mr('，　') +
                 frac(mr('∂') + ssub(mr('U'), mr('2')), mr('∂Q')) + mr('=−') + frac(mr('X'), UN)),
    '2.4': omath(ssub(mr('Q'), mr('PV,max')) + mr('=') +
                 rad(ssup(ssub(mr('S'), mr('N')), mr('2')) + mr('−') +
                     ssup(ssub(mr('P'), mr('PV')), mr('2')))),
    '4.1': omath(ssub(mr('Q'), mr('PV,ref')) + mr('=') + ssub(mr('k'), mr('q')) +
                 dparen(ssub(mr('U'), mr('c')) + mr('−') + ssub(mr('U'), mr('m')))),
    '5.1': omath(ssub(mr('E'), mr('N')) + mr('≥') +
                 frac(ssub(mr('P'), mr('av')) + ssub(mr('t'), mr('max')),
                      ssub(mr('η'), mr('c')) + mr('ΔSOC'))),
}

# ===================== 通用工具 =====================
def set_run_fonts(r, ea='宋体', ascii_f='Times New Roman'):
    r.font.name = ascii_f
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    rf.set(qn('w:ascii'), ascii_f)
    rf.set(qn('w:hAnsi'), ascii_f)
    rf.set(qn('w:eastAsia'), ea)

def style_ea_font(st, ea, ascii_f='Times New Roman'):
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    rf.set(qn('w:ascii'), ascii_f)
    rf.set(qn('w:hAnsi'), ascii_f)
    rf.set(qn('w:eastAsia'), ea)

def first_line_chars(p, n=200):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), str(n))
    ind.set(qn('w:firstLine'), '480' if n == 200 else '0')

INLINE = re.compile(r'_\{([^}]*)\}|\^\{([^}]*)\}')

def emit_inline(p, text, bold=False):
    """解析 _{下标} ^{上标}"""
    pos = 0
    for mobj in INLINE.finditer(text):
        if mobj.start() > pos:
            r = p.add_run(text[pos:mobj.start()])
            set_run_fonts(r); r.bold = bold
        if mobj.group(1) is not None:
            r = p.add_run(mobj.group(1)); set_run_fonts(r); r.bold = bold; r.font.subscript = True
        else:
            r = p.add_run(mobj.group(2)); set_run_fonts(r); r.bold = bold; r.font.superscript = True
        pos = mobj.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        set_run_fonts(r); r.bold = bold

BOLD = re.compile(r'\*\*(.+?)\*\*')
CITE = re.compile(r'\[\[(\d+(?:[-,]\d+)*)\]\]')

def emit_text(p, text):
    """解析 **粗体**、[[n]]上标引用、_{下标}、^{上标}"""
    pos = 0
    marks = []
    for m in BOLD.finditer(text):
        marks.append((m.start(), m.end(), 'b', m.group(1)))
    for m in CITE.finditer(text):
        marks.append((m.start(), m.end(), 'c', m.group(1)))
    marks.sort()
    for s, e, kind, content in marks:
        if s < pos:
            continue
        if s > pos:
            emit_inline(p, text[pos:s])
        if kind == 'b':
            # 粗体段内仍可能有上下标
            emit_inline(p, content, bold=True)
        else:
            r = p.add_run('[' + content + ']')
            set_run_fonts(r); r.font.superscript = True
        pos = e
    if pos < len(text):
        emit_inline(p, text[pos:])

def body_para(doc, text, indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        first_line_chars(p, 200)
    emit_text(p, text)
    return p

def add_formula(doc, key):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.tab_stops.add_tab_stop(Cm(7.9), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run('\t')
    set_run_fonts(r)
    p._p.append(parse_xml(FORMULAS[key]))
    r2 = p.add_run('\t（' + key + '）')
    set_run_fonts(r2)
    return p

def add_figure(doc, img, caption, width_cm):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run()
    run.add_picture(img, width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.line_spacing = 1.0
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run(caption)
    set_run_fonts(r, ea='宋体')
    r.font.size = Pt(10.5)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'bottom', 'left', 'right'):
        if edge in kwargs:
            val = kwargs[edge]
            tag = 'w:{}'.format(edge)
            el = borders.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                borders.append(el)
            el.set(qn('w:val'), val[0])
            el.set(qn('w:sz'), str(val[1]))
            el.set(qn('w:color'), 'auto')

def three_line_table(doc, caption, headers, rows, widths_cm, left_cols=()):
    cap = doc.add_paragraph()
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.line_spacing = 1.0
    cap.paragraph_format.space_before = Pt(6)
    r = cap.add_run(caption)
    set_run_fonts(r, ea='黑体')
    r.bold = True
    r.font.size = Pt(10.5)

    tb = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb.autofit = False
    tblPr = tb._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side, v in (('top', 40), ('left', 80), ('bottom', 40), ('right', 80)):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tblPr.append(mar)

    def fill(cells, values, bold=False):
        for i, v in enumerate(values):
            c = cells[i]
            c.width = Cm(widths_cm[i])
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.alignment = (WD_ALIGN_PARAGRAPH.LEFT if i in left_cols and not bold
                                            else WD_ALIGN_PARAGRAPH.CENTER)
            emit_text(p, v)
            for rr in p.runs:
                if bold:
                    rr.bold = True
                if not rr.font.size:
                    rr.font.size = Pt(10.5)

    fill(tb.rows[0].cells, headers, bold=True)
    for ri, row in enumerate(rows):
        fill(tb.rows[ri + 1].cells, row)
    for ci in range(len(headers)):
        set_cell_border(tb.rows[0].cells[ci], top=('single', 12), bottom=('single', 6))
        set_cell_border(tb.rows[-1].cells[ci], bottom=('single', 12))
    trPr = tb.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
    for row in tb.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
    tail = doc.add_paragraph()
    tail.paragraph_format.line_spacing = 1.0
    tail.paragraph_format.space_after = Pt(6)
    tr = tail.add_run(' ')
    set_run_fonts(tr)
    tr.font.size = Pt(6)

def heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        set_run_fonts(r, ea='黑体')
        r.font.color.rgb = RGBColor(0, 0, 0)
        if level == 1:
            r.font.size = Pt(16); r.bold = True
        elif level == 2:
            r.font.size = Pt(12); r.bold = True
        else:
            r.font.size = Pt(12); r.bold = False
    return h

def setup_styles(doc):
    st = doc.styles['Normal']
    st.font.size = Pt(12)
    style_ea_font(st, '宋体')
    st.paragraph_format.line_spacing = 1.5
    h1 = doc.styles['Heading 1']
    h1.font.size = Pt(16); h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    style_ea_font(h1, '黑体')
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.page_break_before = True
    h1.paragraph_format.line_spacing = 1.5
    h1.paragraph_format.space_before = Pt(18); h1.paragraph_format.space_after = Pt(14)
    h2 = doc.styles['Heading 2']
    h2.font.size = Pt(12); h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    style_ea_font(h2, '黑体')
    h2.paragraph_format.line_spacing = 1.5
    h2.paragraph_format.space_before = Pt(7); h2.paragraph_format.space_after = Pt(7)
    h2.paragraph_format.keep_with_next = True
    h3 = doc.styles['Heading 3']
    h3.font.size = Pt(12); h3.font.bold = False
    h3.font.color.rgb = RGBColor(0, 0, 0)
    style_ea_font(h3, '黑体')
    h3.paragraph_format.line_spacing = 1.5
    h3.paragraph_format.space_before = Pt(7); h3.paragraph_format.space_after = Pt(7)

def strip_doc_grid(section):
    """移除docGrid元素: Word无type不启用网格, LibreOffice却会按网格放大行高, 删除以保证渲染一致"""
    dg = section._sectPr.find(qn('w:docGrid'))
    if dg is not None:
        section._sectPr.remove(dg)

def setup_sections(doc):
    sec1 = doc.sections[0]
    sec1.page_width = Mm(210); sec1.page_height = Mm(297)
    sec1.top_margin = Mm(30); sec1.bottom_margin = Mm(25)
    sec1.left_margin = Mm(30); sec1.right_margin = Mm(20)
    strip_doc_grid(sec1)
    return sec1

def make_body_section(doc):
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.page_width = Mm(210); sec2.page_height = Mm(297)
    sec2.top_margin = Mm(30); sec2.bottom_margin = Mm(25)
    sec2.left_margin = Mm(30); sec2.right_margin = Mm(20)
    sec2.footer.is_linked_to_previous = False
    sectPr = sec2._sectPr
    pg = sectPr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType'); sectPr.append(pg)
    pg.set(qn('w:start'), '1')
    field_xml = (
        '<w:p {}>'
        '<w:pPr><w:jc w:val="right"/></w:pPr>'
        '<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        '<w:sz w:val="18"/></w:rPr><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        '<w:sz w:val="18"/></w:rPr><w:instrText xml:space="preserve"> PAGE \\* arabic \\* MERGEFORMAT </w:instrText></w:r>'
        '<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>'
        '<w:sz w:val="18"/></w:rPr><w:fldChar w:fldCharType="end"/></w:r>'
        '</w:p>').format(W_DECL)
    fp = sec2.footer.paragraphs[0]
    fp._p.getparent().replace(fp._p, parse_xml(field_xml))
    strip_doc_grid(sec2)
    return sec2

def plain_center(doc, text, ea='黑体', size=16, bold=True, ascii_f='Times New Roman', space_after=14):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_run_fonts(r, ea=ea, ascii_f=ascii_f)
    r.bold = bold
    r.font.size = Pt(size)
    return p

def page_break(doc):
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_break(WD_BREAK.PAGE)
