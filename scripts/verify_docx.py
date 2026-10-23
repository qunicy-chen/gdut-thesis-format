# -*- coding: utf-8 -*-
"""docx 程序化校验: 上标引用完整性 / OMML公式 / 加粗标记 / 图表数量"""
import re, sys, zipfile
import docx
from docx.oxml.ns import qn

DOCX = r'C:\Users\<你的用户名>\Desktop\论文项目\交付\智能配电网光储协同的电压治理系统方案设计_课程报告.docx'

def check():
    # ---- 1. OMML 公式检查 ----
    with zipfile.ZipFile(DOCX) as z:
        xml = z.read('word/document.xml').decode('utf-8')
    n_math = xml.count('<m:oMath')
    print('OMML <m:oMath> 公式数量:', n_math)
    # 图片(嵌入)
    n_img = xml.count('<w:drawing')
    print('嵌入图片数量:', n_img)

    # ---- 2. 上标引用检查 ----
    d = docx.Document(DOCX)
    cites = []
    for p in d.paragraphs:
        for r in p.runs:
            if r.font.superscript:
                m = re.fullmatch(r'\[(\d+(?:[-,]\d+)*)\]', r.text.strip())
                if m:
                    for part in m.group(1).split(','):
                        if '-' in part:
                            a, b = part.split('-')
                            cites += list(range(int(a), int(b) + 1))
                        else:
                            cites.append(int(part))
    print('上标引用出现次数:', len(cites))
    print('被引编号排序:', sorted(set(cites)))
    n_refs = sum(1 for p in d.paragraphs if re.match(r'^\[\d+\] ', p.text.strip()))
    print('参考文献条目数:', n_refs)
    missing = [i for i in range(1, n_refs + 1) if i not in cites]
    excess = [i for i in set(cites) if i > n_refs or i < 1]
    print('未被引用的编号:', missing)
    print('超出范围的编号:', excess)

    # ---- 3. 加粗关键结论 ----
    n_bold = sum(1 for p in d.paragraphs for r in p.runs if r.bold and len(r.text) > 10)
    print('长加粗run数量(关键结论):', n_bold)

    # ---- 4. 图题/表题 ----
    caps = [p.text.strip() for p in d.paragraphs
            if re.match(r'^(图|表|附表)\S+\s+\S', p.text.strip())]
    print('图表题共', len(caps), '条:')
    for c in caps:
        print('  ', c)

    ok = (n_math == 6 and not missing and not excess and n_refs >= 20 and n_img >= 4)
    print('VERIFY:', 'PASS' if ok else 'CHECK-NEEDED')
    return ok

if __name__ == '__main__':
    check()
